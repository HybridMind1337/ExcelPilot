"""Unit tests for SecurityAuditor."""

from pathlib import Path

import openpyxl
import pytest

from excelpilot.exceptions.base import WorkbookError
from excelpilot.security.auditor import SecurityAuditor
from excelpilot.security.models import SecuritySeverity, ThreatCategory


def test_security_clean_workbook(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Clean"
    ws.append(["ID", "Name", "Score"])
    ws.append([1, "Alpha", 100])
    wb_path = tmp_path / "clean.xlsx"
    wb.save(str(wb_path))
    wb.close()

    auditor = SecurityAuditor()
    report = auditor.audit_workbook(wb_path)

    assert report.security_score == 100.0
    assert report.risk_level == "SAFE"
    assert len(report.findings) == 0
    assert report.hidden_sheets_count == 0


def test_security_threats_detection(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Visible"

    # Add dangerous formula injection
    ws1.append(["Payload", "Data"])
    ws1.append(["=cmd|'/C calc'!A0", "safe_val"])

    # Add external reference
    ws1.append(["=[ExternalBook.xlsx]Sheet1!A1", "other"])

    # Add PII
    ws1.append(["test.user@company.co.uk", "4111-2222-3333-4444"])

    # Add hidden worksheet
    ws2 = wb.create_sheet("HiddenSheet")
    ws2.sheet_state = "hidden"
    ws2.append(["Confidential", 12345])

    wb_path = tmp_path / "vulnerable.xlsx"
    wb.save(str(wb_path))
    wb.close()

    auditor = SecurityAuditor()
    report = auditor.audit_workbook(wb_path)

    assert report.hidden_sheets_count == 1
    assert report.security_score < 75.0

    categories = {f.category for f in report.findings}
    assert ThreatCategory.FORMULA_INJECTION in categories
    assert ThreatCategory.EXTERNAL_LINK in categories
    assert ThreatCategory.HIDDEN_SHEET in categories
    assert ThreatCategory.PII_EXPOSURE in categories

    # Verify critical severity for formula injection
    inj = next(f for f in report.findings if f.category == ThreatCategory.FORMULA_INJECTION)
    assert inj.severity == SecuritySeverity.CRITICAL


def test_security_macro_detection(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb_path = tmp_path / "macro_file.xlsm"
    wb.save(str(wb_path))
    wb.close()

    auditor = SecurityAuditor()
    report = auditor.audit_workbook(wb_path)
    assert report.is_macro_enabled is True
    assert any(f.category == ThreatCategory.MACRO_VBA for f in report.findings)


def test_security_missing_file(tmp_path: Path) -> None:
    auditor = SecurityAuditor()
    with pytest.raises(WorkbookError, match="Workbook not found"):
        auditor.audit_workbook(tmp_path / "non_existent.xlsx")
