"""Synthetic workbook generation for automated test fixtures."""

import datetime
from pathlib import Path

import openpyxl


def create_synthetic_sales_workbook(path: Path) -> Path:
    """Create a synthetic sales workbook with realistic transactional records."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sales"

    headers = [
        "OrderID",
        "CustomerID",
        "CustomerName",
        "OrderDate",
        "Region",
        "Channel",
        "Product",
        "Revenue",
        "Cost",
        "Quantity",
        "Email",
        "Phone",
    ]
    ws.append(headers)

    regions = ["North", "South", "East", "West"]
    channels = ["Online", "Retail", "Wholesale"]
    products = [
        "Laptop Pro",
        "Wireless Mouse",
        "USB-C Hub",
        "27-inch Monitor",
        "Ergonomic Keyboard",
    ]

    base_date = datetime.date(2025, 1, 1)

    for i in range(1, 151):
        order_id = f"ORD-{1000 + i}"
        cust_id = f"CUST-{(i % 30) + 1:03d}"
        cust_name = f"Customer {i % 30}"
        order_date = (base_date + datetime.timedelta(days=i % 60)).isoformat()
        region = regions[i % len(regions)]
        channel = channels[i % len(channels)]
        product = products[i % len(products)]
        quantity = (i % 5) + 1
        unit_price = 45.0 + (i % 10) * 15.0
        revenue = round(quantity * unit_price, 2)
        cost = round(revenue * 0.65, 2)
        email = f"user_{i % 30}@example.com"
        phone = f"+44 7700 900{i % 90:02d}"

        ws.append(
            [
                order_id,
                cust_id,
                cust_name,
                order_date,
                region,
                channel,
                product,
                revenue,
                cost,
                quantity,
                email,
                phone,
            ]
        )

    wb.save(str(path))
    wb.close()
    return path


def create_synthetic_quality_issue_workbook(path: Path) -> Path:
    """Create a synthetic workbook containing deliberate quality defects."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Transactions"

    headers = ["OrderID", "CustomerID", "Amount", "Quantity", "Category"]
    ws.append(headers)

    # Valid rows
    ws.append(["ORD-001", "C-01", 100.0, 2, "Hardware"])
    ws.append(["ORD-002", "C-02", 250.0, 5, "Software"])

    # Duplicate row
    ws.append(["ORD-002", "C-02", 250.0, 5, "Software"])

    # Negative quantity (validity defect)
    ws.append(["ORD-003", "C-03", 50.0, -3, "Hardware"])

    # Missing customer ID (completeness defect)
    ws.append(["ORD-004", None, 120.0, 1, "Hardware"])

    # Untrimmed whitespace (consistency defect)
    ws.append(["ORD-005", "C-05", 90.0, 1, "  Hardware  "])

    wb.save(str(path))
    wb.close()
    return path


def create_empty_worksheet_workbook(path: Path) -> Path:
    """Create a workbook containing an empty sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "EmptySheet"
    wb.save(str(path))
    wb.close()
    return path
